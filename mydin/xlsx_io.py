"""Backup e restauração em .xlsx — abas separadas, restauração fiel (valores brutos)."""
import datetime as dt
import io

from openpyxl import Workbook, load_workbook

TABELAS = [
    ("transacoes", ["id", "data", "valor_cent", "tipo", "categoria_id", "descricao", "origem",
                    "id_externo", "conta", "contato_id", "projeto_id", "status_revisao",
                    "especial", "compartilhada", "compra_pai_id", "meia"]),
    ("contatos", ["id", "nome", "cpf_mascarado", "tipo", "aluno_id", "sempre_revisar"]),
    ("alunos", ["id", "nome", "turma", "ativo", "valor_cota_cent"]),
    ("projetos", ["id", "cliente_id", "titulo", "descricao", "modalidade", "valor_orcado_cent",
                  "status", "valor_sinal_cent", "sinal_pago", "valor_saldo_cent", "saldo_pago",
                  "criado_em", "concluido_em"]),
    ("registros_tempo", ["id", "projeto_id", "data", "horas", "descricao", "metodo"]),
    ("pendencias", ["id", "tipo", "contato_id", "projeto_id", "valor_esperado_cent",
                    "data_esperada", "status", "transacao_id", "obs"]),
    ("categorias", ["id", "nome", "grupo"]),
    ("regras", ["id", "tipo", "contato_id", "padrao", "categoria_id", "criado_em"]),
    ("vendedores_med", ["id", "contato_id"]),
    ("config", ["chave", "valor"]),
]

# Ordem de inserção que respeita as FKs; deleção na ordem inversa
ORDEM_RESTORE = ["categorias", "alunos", "contatos", "projetos", "config",
                 "transacoes", "registros_tempo", "pendencias", "regras", "vendedores_med"]


def nome_backup():
    return "mydin_backup_%s.xlsx" % dt.datetime.now().strftime("%Y%m%d_%H%M%S")


def exportar(db):
    wb = Workbook()
    wb.remove(wb.active)
    for tabela, colunas in TABELAS:
        ws = wb.create_sheet(tabela)
        ws.append(colunas)
        for row in db.execute(f"SELECT {', '.join(colunas)} FROM {tabela}"):
            ws.append([row[c] for c in colunas])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def restaurar(db, arquivo_bytes):
    """Substitui TODO o banco pelo conteúdo do backup. Retorna contagens por tabela."""
    wb = load_workbook(io.BytesIO(arquivo_bytes), data_only=True)
    colunas_por_tabela = dict(TABELAS)
    faltando = [t for t in ORDEM_RESTORE if t not in wb.sheetnames]
    if faltando:
        raise ValueError("Backup inválido — abas ausentes: " + ", ".join(faltando))

    db.execute("PRAGMA foreign_keys = OFF")
    try:
        for tabela in reversed(ORDEM_RESTORE):
            db.execute(f"DELETE FROM {tabela}")
        contagens = {}
        for tabela in ORDEM_RESTORE:
            colunas = colunas_por_tabela[tabela]
            ws = wb[tabela]
            linhas = list(ws.iter_rows(min_row=1, values_only=True))
            if not linhas:
                contagens[tabela] = 0
                continue
            header = [str(h) if h is not None else "" for h in linhas[0]]
            idx = {c: header.index(c) for c in colunas if c in header}
            n = 0
            for linha in linhas[1:]:
                if linha is None or all(v is None for v in linha):
                    continue
                valores = [linha[idx[c]] if c in idx else None for c in colunas]
                db.execute(
                    f"INSERT INTO {tabela} ({', '.join(colunas)}) VALUES ({', '.join('?' * len(colunas))})",
                    valores,
                )
                n += 1
            contagens[tabela] = n
        db.commit()
    finally:
        db.execute("PRAGMA foreign_keys = ON")
    return contagens
